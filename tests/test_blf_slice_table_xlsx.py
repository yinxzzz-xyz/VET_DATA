import re
import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
import openpyxl
from vet_data_modular.blf_slice_table import TableParseError, parse_condition_table

class XlsxConditionTableTests(unittest.TestCase):
    def _save(self, folder, rows, second_sheet_rows=None):
        path = Path(folder) / "conditions.xlsx"
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "工况表"
        for row in rows:
            sheet.append(row)
        if second_sheet_rows is not None:
            other = book.create_sheet("不应读取")
            for row in second_sheet_rows:
                other.append(row)
        book.save(path)
        book.close()
        return path

    def _add_formula_cache(self, path, cell_reference, cached_excel_value):
        replacement = Path(str(path) + ".replacement")
        with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(replacement, "w") as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    xml = data.decode("utf-8")
                    pattern = rf'(<c r="{cell_reference}"[^>]*>.*?<f[^>]*>.*?</f>)<v></v>'
                    xml, count = re.subn(pattern, rf"\1<v>{cached_excel_value}</v>", xml, count=1)
                    self.assertEqual(count, 1)
                    data = xml.encode("utf-8")
                target.writestr(item, data)
        replacement.replace(path)

    def test_reads_only_first_sheet_and_native_datetime(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "记录内容"], [datetime(2026, 6, 29, 14, 30, 15), "中文工况"]], [["记录时间", "记录内容"], [datetime(2026, 1, 1), "第二页"]])
            result = parse_condition_table(path)
        self.assertEqual(result.worksheet_name, "工况表")
        self.assertIsNone(result.encoding)
        self.assertEqual([item.name for item in result.conditions], ["中文工况"])
        self.assertEqual(result.conditions[0].recorded_at, datetime(2026, 6, 29, 14, 30, 15))

    def test_text_formats_and_minute_precision(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "记录内容"], ["2026-06-29 14:30:15", "秒格式"], ["2026/6/29 9:05", "分钟格式"]])
            result = parse_condition_table(path)
        self.assertEqual(result.conditions[0].recorded_at.second, 15)
        self.assertEqual(result.conditions[1].recorded_at, datetime(2026, 6, 29, 9, 5, 0))

    def test_additional_columns_and_optional_seconds(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["序号", "记录时间", "记录内容", "纬度", "向前秒数", "向后秒数"], [1, "2026/6/29 9:05", "自定义", 31.2, 0, 125], [2, "2026/6/29 9:06", "默认", 31.3, None, None]])
            result = parse_condition_table(path)
        self.assertEqual((result.conditions[0].before_seconds, result.conditions[0].after_seconds), (0, 125))
        self.assertEqual((result.conditions[1].before_seconds, result.conditions[1].after_seconds), (60, 60))

    def test_missing_required_header_is_file_error(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "其他"], ["2026/6/29 9:05", "值"]])
            with self.assertRaisesRegex(TableParseError, "缺少必需表头: 记录内容"):
                parse_condition_table(path)

    def test_invalid_rows_boundaries_and_original_row_numbers(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "记录内容"], [None, None], [45200, "普通数字"], ["not-a-date", "无效日期"], ["1999-12-31 23:59:59", "越界"], ["2000-01-01 00:00:00", "最小边界"], ["2100-12-31 23:59:59", "最大边界"], ["2101-01-01 00:00:00", "越界"], ["2026/6/29 9:05", "有效"]])
            result = parse_condition_table(path)
        self.assertEqual([item.name for item in result.conditions], ["最小边界", "最大边界", "有效"])
        self.assertEqual([item.original_row_number for item in result.conditions], [6, 7, 9])
        self.assertEqual(len(result.discarded_rows), 5)
        self.assertEqual(result.discarded_rows[0].source_location, "工作表“工况表”第 2 行")
        self.assertIn("不能使用普通数字", result.discarded_rows[1].reason)

    def test_formula_with_valid_cached_datetime_is_used(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "记录内容"], ["=DATE(2026,6,29)+TIME(9,5,0)", "缓存公式"]])
            book = openpyxl.load_workbook(path)
            book.active["A2"].number_format = "yyyy-mm-dd hh:mm:ss"
            book.save(path)
            book.close()
            self._add_formula_cache(path, "A2", "46202.37847222222")
            result = parse_condition_table(path)
        self.assertEqual(result.conditions[0].recorded_at, datetime(2026, 6, 29, 9, 5, 0))

    def test_formula_without_cache_is_discarded_and_valid_row_survives(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "记录内容"], ["=NOW()", "无缓存"], ["2026/6/29 9:05", "有效"]])
            result = parse_condition_table(path)
        self.assertEqual([item.name for item in result.conditions], ["有效"])
        self.assertEqual(result.discarded_rows[0].row_number, 2)
        self.assertEqual(result.discarded_rows[0].worksheet_name, "工况表")
        self.assertIn("公式没有有效缓存结果", result.discarded_rows[0].reason)

    def test_no_valid_conditions_is_explicit(self):
        with tempfile.TemporaryDirectory() as folder:
            path = self._save(folder, [["记录时间", "记录内容"], [None, None], [123, "数字"]])
            result = parse_condition_table(path)
        self.assertEqual(result.conditions, ())
        self.assertEqual(len(result.discarded_rows), 2)

if __name__ == "__main__":
    unittest.main()
