"""CSV and XLSX condition-table parsing for BLF slicing."""
from __future__ import annotations
import csv
import re
from dataclasses import dataclass, field
from datetime import datetime
from io import StringIO
from pathlib import Path
from typing import Any, Mapping, Sequence
import openpyxl
from .blf_slice_models import (DEFAULT_AFTER_SECONDS, DEFAULT_BEFORE_SECONDS, HEADER_MAX_YEAR, HEADER_MIN_YEAR, ConditionSpec)

REQUIRED_HEADERS = ("记录时间", "记录内容")
OPTIONAL_BEFORE_HEADER = "向前秒数"
OPTIONAL_AFTER_HEADER = "向后秒数"
SUPPORTED_TIME_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M")
_INTEGER_PATTERN = re.compile(r"[+-]?\d+\Z")

class TableParseError(ValueError):
    """A file-level error that prevents the condition table from being used."""

@dataclass(frozen=True, slots=True)
class DiscardedRow:
    """One rejected table record and its source location."""
    row_number: int
    end_row_number: int
    reason: str
    values: tuple[Any, ...] = field(default_factory=tuple)
    worksheet_name: str | None = None

    @property
    def source_location(self) -> str:
        prefix = f"工作表“{self.worksheet_name}”" if self.worksheet_name else ""
        rows = f"第 {self.row_number} 行" if self.row_number == self.end_row_number else f"第 {self.row_number}-{self.end_row_number} 行"
        return prefix + rows

@dataclass(frozen=True, slots=True)
class TableParseResult:
    """Parsed conditions and diagnostics needed by confirmation/report layers."""
    source_path: Path
    encoding: str | None
    conditions: tuple[ConditionSpec, ...] = field(default_factory=tuple)
    discarded_rows: tuple[DiscardedRow, ...] = field(default_factory=tuple)
    worksheet_name: str | None = None

    def __post_init__(self) -> None:
        object.__setattr__(self, "source_path", Path(self.source_path))
        object.__setattr__(self, "conditions", tuple(self.conditions))
        object.__setattr__(self, "discarded_rows", tuple(self.discarded_rows))

def _decode_csv(path: Path) -> tuple[str, str]:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise TableParseError(f"无法读取工况表: {exc}") from exc
    try:
        return raw.decode("utf-8-sig"), "utf-8-sig"
    except UnicodeDecodeError:
        try:
            return raw.decode("gbk"), "gbk"
        except UnicodeDecodeError as exc:
            raise TableParseError("工况表既不是有效的 UTF-8，也不是有效的 GBK 编码。") from exc

def _parse_recorded_at(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        text = value.strip()
        if not text:
            raise ValueError("记录时间为空")
        for fmt in SUPPORTED_TIME_FORMATS:
            try:
                parsed = datetime.strptime(text, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError("记录时间格式无效")
    elif value is None:
        raise ValueError("记录时间为空")
    else:
        raise ValueError("记录时间必须为 Excel 日期时间或支持的文本格式，不能使用普通数字")
    if not HEADER_MIN_YEAR <= parsed.year <= HEADER_MAX_YEAR:
        raise ValueError(f"记录时间年份必须在 {HEADER_MIN_YEAR}～{HEADER_MAX_YEAR} 之间")
    return parsed

def _parse_optional_seconds(value: Any, default: int, header: str) -> int:
    if value is None or isinstance(value, str) and not value.strip():
        return default
    if isinstance(value, bool):
        raise ValueError(f"{header}必须为整数")
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = value.strip() if isinstance(value, str) else ""
    if not _INTEGER_PATTERN.fullmatch(text):
        raise ValueError(f"{header}必须为整数")
    return int(text)

def _validate_headers(headers: Sequence[Any]) -> dict[str, int]:
    missing = [h for h in REQUIRED_HEADERS if h not in headers]
    duplicated = [h for h in REQUIRED_HEADERS if headers.count(h) > 1]
    if missing:
        raise TableParseError(f"缺少必需表头: {', '.join(missing)}")
    if duplicated:
        raise TableParseError(f"必需表头重复: {', '.join(duplicated)}")
    columns = {h: headers.index(h) for h in REQUIRED_HEADERS}
    for header in (OPTIONAL_BEFORE_HEADER, OPTIONAL_AFTER_HEADER):
        if header in headers:
            columns[header] = headers.index(header)
    return columns

def _build_condition(values: Mapping[str, Any], row_number: int) -> ConditionSpec:
    value = values.get("记录内容")
    name = value if isinstance(value, str) else str(value or "")
    if not name.strip():
        raise ValueError("记录内容为空")
    return ConditionSpec(row_number, name, name, _parse_recorded_at(values.get("记录时间")), _parse_optional_seconds(values.get(OPTIONAL_BEFORE_HEADER), DEFAULT_BEFORE_SECONDS, OPTIONAL_BEFORE_HEADER), _parse_optional_seconds(values.get(OPTIONAL_AFTER_HEADER), DEFAULT_AFTER_SECONDS, OPTIONAL_AFTER_HEADER))

def _parse_csv(path: Path) -> TableParseResult:
    text, encoding = _decode_csv(path)
    reader = csv.reader(StringIO(text, newline=""), strict=True)
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise TableParseError("工况表为空，缺少表头。") from exc
    except csv.Error as exc:
        raise TableParseError(f"CSV 表头解析失败: {exc}") from exc
    columns = _validate_headers(headers)
    conditions, discarded = [], []
    next_row = reader.line_num + 1
    try:
        for row in reader:
            row_number, end_row = next_row, reader.line_num
            next_row = end_row + 1
            if not row or not any(v.strip() for v in row):
                discarded.append(DiscardedRow(row_number, end_row, "空白行", tuple(row)))
                continue
            selected = {h: row[i] if i < len(row) else "" for h, i in columns.items()}
            try:
                conditions.append(_build_condition(selected, row_number))
            except ValueError as exc:
                discarded.append(DiscardedRow(row_number, end_row, str(exc), tuple(row)))
    except csv.Error as exc:
        raise TableParseError(f"CSV 在第 {reader.line_num} 行附近解析失败: {exc}") from exc
    return TableParseResult(path, encoding, tuple(conditions), tuple(discarded))

def _parse_xlsx(path: Path) -> TableParseResult:
    try:
        formula_book = openpyxl.load_workbook(path, read_only=False, data_only=False)
        cached_book = openpyxl.load_workbook(path, read_only=False, data_only=True)
    except (OSError, ValueError, KeyError, openpyxl.utils.exceptions.InvalidFileException) as exc:
        raise TableParseError(f"无法读取 XLSX 工况表: {exc}") from exc
    try:
        formula_sheet, cached_sheet = formula_book.worksheets[0], cached_book.worksheets[0]
        sheet_name = formula_sheet.title
        formula_rows, cached_rows = formula_sheet.iter_rows(), cached_sheet.iter_rows()
        try:
            headers = [cell.value for cell in next(formula_rows)]
            next(cached_rows)
        except StopIteration as exc:
            raise TableParseError(f"工作表“{sheet_name}”为空，缺少表头。") from exc
        columns = _validate_headers(headers)
        conditions, discarded = [], []
        for row_number, (raw_row, cache_row) in enumerate(zip(formula_rows, cached_rows), start=2):
            raw_values = tuple(cell.value for cell in raw_row)
            if not any(v is not None and (not isinstance(v, str) or v.strip()) for v in raw_values):
                discarded.append(DiscardedRow(row_number, row_number, "空白行", raw_values, sheet_name))
                continue
            selected, formula_error = {}, None
            for header, index in columns.items():
                if index >= len(raw_row):
                    selected[header] = None
                elif raw_row[index].data_type == "f":
                    selected[header] = cache_row[index].value
                    if selected[header] is None:
                        formula_error = f"{header}公式没有有效缓存结果"
                        break
                else:
                    selected[header] = raw_row[index].value
            try:
                if formula_error:
                    raise ValueError(formula_error)
                conditions.append(_build_condition(selected, row_number))
            except ValueError as exc:
                discarded.append(DiscardedRow(row_number, row_number, str(exc), raw_values, sheet_name))
        return TableParseResult(path, None, tuple(conditions), tuple(discarded), sheet_name)
    finally:
        formula_book.close()
        cached_book.close()

def parse_condition_table(path: str | Path) -> TableParseResult:
    """Parse a CSV or XLSX condition table without task-level validation."""
    source_path = Path(path)
    if source_path.suffix.lower() == ".csv":
        return _parse_csv(source_path)
    if source_path.suffix.lower() == ".xlsx":
        return _parse_xlsx(source_path)
    raise TableParseError(f"不支持的工况表格式: {source_path.suffix or '(无扩展名)'}")
